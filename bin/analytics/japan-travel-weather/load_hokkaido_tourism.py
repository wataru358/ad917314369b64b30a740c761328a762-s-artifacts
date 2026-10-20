import hashlib
import math
import re
import unicodedata
from pathlib import Path

import duckdb
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


BASE_DIR = Path(__file__).resolve().parent
INPUT_DIR = BASE_DIR / "data" / "tourism" / "hokkaido"
OUTPUT_DIR = BASE_DIR / "data" / "processed"

SHEET = "6～28頁"
MONTHS = [4, 5, 6, 7, 8, 9, 10, 11, 12, 1, 2, 3]
LABELS = [
    "入込総数", "内道外客", "内道内客",
    "内日帰客", "内宿泊客", "宿泊客延数",
]
OFFSETS = [0, 3, 4, 5]  # 入込総数の行から数えた4項目の位置。


def clean(value):
    # 全角数字を半角にし、見出しや市町村名の空白を除く。
    text = "" if value is None else str(value)
    return re.sub(r"\s+", "", unicodedata.normalize("NFKC", text))


def number(value, location):
    # 空欄・文字列・負数を0に置き換えず、原表の確認を求める。
    if (
        isinstance(value, bool)
        or not isinstance(value, (int, float))
        or not math.isfinite(value)
        or value < 0
    ):
        raise ValueError(
            f"数値を確認してください: {location}: {value!r}"
        )
    return float(value)


records = []
reference_names = None
annual_checks = 0

for fiscal_year in range(2022, 2026):
    extension = "xlsx" if fiscal_year <= 2023 else "xlsm"
    path = INPUT_DIR / (
        f"hokkaido_tourism_fy{fiscal_year}_tables.{extension}"
    )
    source_hash = hashlib.sha256(path.read_bytes()).hexdigest()

    # Excelに保存済みの計算結果を読む。原本は保存し直さない。
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = list(wb[SHEET].iter_rows(values_only=True))

        title = clean(rows[0][0])
        match = re.search(r"令和(\d+)年度", title)
        if not match or int(match.group(1)) + 2018 != fiscal_year:
            raise ValueError(
                f"対象年度が一致しません: {path.name}: {title}"
            )
        if "市町村別・月別観光入込客数" not in title:
            raise ValueError(f"表の種類が違います: {path.name}")
        if (
            "入込総数→千人、宿泊客延数→千人泊"
            not in clean(rows[1][18])
        ):
            raise ValueError(f"単位が想定と異なります: {path.name}")

        headers = [clean(v) for v in rows[2][2:17]]
        expected_headers = (
            ["市町村名", "区分"]
            + [f"{m}月" for m in MONTHS]
            + ["合計"]
        )
        if headers != expected_headers:
            raise ValueError(f"列の構成が違います: {path.name}")

        # C列が市町村名、D列が「入込総数」の行を開始位置にする。
        # 全道計・圏域計・振興局計はC列が空なので取り込まない。
        starts = [
            i for i, row in enumerate(rows)
            if clean(row[2]) and clean(row[3]) == "入込総数"
        ]
        names = [clean(rows[i][2]) for i in starts]

        if len(names) != 179 or len(set(names)) != 179:
            raise ValueError(
                f"179市町村が一意にそろっていません: {path.name}"
            )
        if reference_names is None:
            reference_names = set(names)
        elif set(names) != reference_names:
            raise ValueError(
                f"年度間で市町村が一致しません: {path.name}"
            )

        before = len(records)

        for i in starts:
            municipality = clean(rows[i][2])
            actual_labels = [
                clean(rows[i + offset][3])
                for offset in range(6)
            ]
            if actual_labels != LABELS:
                raise ValueError(
                    f"6行の区分が想定と違います: {municipality}"
                )

            values = []
            for offset in OFFSETS:
                row = rows[i + offset]
                location = (
                    f"{path.name}/{SHEET}/行{i + offset + 1}"
                )
                monthly = [
                    number(v, location) for v in row[4:16]
                ]
                annual = number(row[16], location)

                # 暦年で切り出す前に、4〜3月の合計をQ列と照合する。
                if not math.isclose(
                    sum(monthly), annual,
                    rel_tol=0, abs_tol=1e-6,
                ):
                    raise ValueError(
                        f"12か月合計とQ列が一致しません: {location}"
                    )

                annual_checks += 1
                values.append(monthly)

            for j, month in enumerate(MONTHS):
                # 1〜3月は、年度の翌年に属する。
                year = (
                    fiscal_year if month >= 4
                    else fiscal_year + 1
                )
                if not 2023 <= year <= 2025:
                    continue

                total, day, overnight, nights = [
                    v[j] for v in values
                ]
                records.append((
                    year, month, municipality,
                    total, day, overnight, nights,
                    fiscal_year, path.name, SHEET, i + 1,
                    get_column_letter(j + 5), source_hash,
                ))

        print(
            f"{fiscal_year}年度: 179市町村 / "
            f"採用 {len(records) - before:,}行"
        )
    finally:
        wb.close()

# 市町村×暦年×月が、ちょうど179×3×12通りあるか確認する。
expected_keys = {
    (year, month, name)
    for year in (2023, 2024, 2025)
    for month in range(1, 13)
    for name in reference_names
}
actual_keys = {(r[0], r[1], r[2]) for r in records}

if (
    len(records) != len(expected_keys)
    or actual_keys != expected_keys
):
    raise ValueError("年月・市町村に不足、余分、重複があります。")

con = duckdb.connect()
try:
    # 元の単位を列名に明記する。小数値を保つためDOUBLEで保存する。
    con.execute("""
        CREATE TABLE tourism (
            year INTEGER,
            month INTEGER,
            municipality VARCHAR,
            visitors_thousands DOUBLE,
            day_visitors_thousands DOUBLE,
            overnight_visitors_thousands DOUBLE,
            guest_nights_thousands DOUBLE,
            fiscal_year INTEGER,
            source_file VARCHAR,
            source_sheet VARCHAR,
            source_row INTEGER,
            source_column VARCHAR,
            source_sha256 VARCHAR
        )
    """)
    con.executemany(
        "INSERT INTO tourism VALUES "
        "(?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
        records,
    )

    # 今回使う2023〜2025年で、総数と日帰り＋宿泊を照合する。
    differences = con.sql("""
        SELECT year, month, municipality
        FROM tourism
        WHERE abs(
            visitors_thousands
            - day_visitors_thousands
            - overnight_visitors_thousands
        ) > 1e-6
    """).fetchall()

    if differences:
        raise ValueError(
            f"内訳の差を原表で確認してください: {differences[:5]}"
        )

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    stem = "hokkaido_tourism_monthly_2023_2025"

    result = con.sql("""
        SELECT * FROM tourism
        ORDER BY municipality, year, month
    """)
    result.write_parquet(
        str(OUTPUT_DIR / f"{stem}.parquet"),
        compression="zstd",
    )
    result.write_csv(
        str(OUTPUT_DIR / f"{stem}.csv"),
        header=True,
    )

    # 保存後のParquetを読み直し、全列・全行の一致を確認する。
    con.read_parquet(
        str(OUTPUT_DIR / f"{stem}.parquet")
    ).create_view("saved")

    difference = con.sql("""
        (SELECT * FROM tourism EXCEPT ALL SELECT * FROM saved)
        UNION ALL
        (SELECT * FROM saved EXCEPT ALL SELECT * FROM tourism)
    """).fetchmany(1)

    if difference:
        raise ValueError("Parquetの保存前後で内容が一致しません。")

    print(
        f"検査OK: 年度合計 {annual_checks:,}件、"
        "月別内訳 6,444件、一致"
    )
    print(
        "検査OK: 179市町村×36か月、"
        "重複・欠損なし、Parquet全列一致"
    )
    print(f"保存: {stem}.parquet / {len(records):,}行")
    print(f"保存: {stem}.csv / {len(records):,}行")

    con.sql("""
        SELECT year,
            count(*) AS rows,
            count(DISTINCT municipality) AS municipalities,
            count(DISTINCT month) AS months
        FROM tourism
        GROUP BY year
        ORDER BY year
    """).show()

finally:
    con.close()
