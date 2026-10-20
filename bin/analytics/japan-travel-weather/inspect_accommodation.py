from pathlib import Path
from openpyxl import load_workbook


BASE_DIR = Path(__file__).resolve().parent
INPUT_DIR = BASE_DIR / "data" / "accommodation"

files = sorted(INPUT_DIR.glob("*.xlsx"))

if not files:
    raise FileNotFoundError(f"Excelが見つかりません: {INPUT_DIR}")

for path in files:
    print(f"\nfile: {path.name}")

    wb = load_workbook(path, read_only=True, data_only=True)

    # try:
    #     print("sheet list: ", wb.sheetnames)

    #     sheet = wb["第2表(1月)"]
    #     print(f"\nsheet to confirm: ", {sheet.title})

    #     # for row_number, row in enumerate(
    #     #     sheet.iter_rows(max_row=20, max_col=16, values_only=True),
    #     #     start=1,
    #     # ):
    #     #     print(row_number, row)
    #     for area, guest_nights in sheet.iter_rows(
    #         min_row=8,
    #         max_row=54,
    #         max_col=2,
    #         values_only=True,
    #     ):
    #         area = area.strip()

    #         prefecture_code = area[:2]
    #         prefecture = area[2:]

    #         print(prefecture_code, prefecture, guest_nights)
    try:
        # range(1, 13)は、1〜12を順番に返す。13は含まない。
        for month in range(1, 13):
            sheet = wb[f"第2表({month}月)"]

            # B列だけを読むと対象年が分からなくなるので、
            # A7にある年月も表示する。
            print(f"\n対象年月: {sheet['A7'].value}")

            for area, guest_nights in sheet.iter_rows(
                min_row=8,
                max_row=54,
                max_col=2,
                values_only=True,
            ):
                area = area.strip()
                prefecture_code = area[:2]
                prefecture = area[2:]

                print(month, prefecture_code, prefecture, guest_nights)
    finally:
        wb.close()
