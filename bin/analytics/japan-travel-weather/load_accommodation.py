import re
from pathlib import Path

import duckdb
from openpyxl import load_workbook


BASE_DIR = Path(__file__).resolve().parent
INPUT_DIR = BASE_DIR / "data" / "accommodation"
OUTPUT_FILE = (
    BASE_DIR / "data" / "processed"
    / "accommodation_monthly_2023_2025.parquet"
)

INPUT_FILES = {
    "001750679.xlsx": 2023,
    "001905499.xlsx": 2024,
    "002010340.xlsx": 2025,
}

records = []

for filename, year in INPUT_FILES.items():
    path = INPUT_DIR / filename
    if not path.is_file():
        raise FileNotFoundError(f"no files: {path}")

    wb = load_workbook(path, read_only=True, data_only=True)

    try:
        for month in range(1, 13):
            sheet = wb[f"第2表({month}月)"]
            # 例：「令和5年1月」から、年と月の数字を取り出す。
            label = str(sheet["A7"].value).strip()
            match = re.fullmatch(r"令和(\d+)年(\d+)月", label)

            # 令和の年数に2018を足すと西暦になる。
            if not match or (
                int(match.group(1)) + 2018,
                int(match.group(2)),
            ) != (year, month):
                raise ValueError(
                    f"{filename}/{sheet.title}: 対象年月が違います: {label}"
                )

            for area, guest_nights in sheet.iter_rows(
                min_row=8,
                max_row=54,
                max_col=2,
                values_only=True,
            ):
                # 都道府県コードと名前を分ける。
                area = str(area).strip()
                area_match = re.fullmatch(r"(\d{2})(.+)", area)
                if not area_match:
                    raise ValueError(f"都道府県の表記が不正です: {area}")

                prefecture_code = area_match.group(1)
                prefecture = area_match.group(2)

                # 空欄や「-」を、宿泊者数0として扱わない。
                # 今回は数値でない値があれば、保存前に止める。
                if (
                    not isinstance(guest_nights, (int, float))
                    or guest_nights < 0
                    or not float(guest_nights).is_integer()
                ):
                    raise ValueError(
                        f"{year}年{month}月/{prefecture}: "
                        f"宿泊者数を確認してください: {guest_nights}"
                    )

                # 元のファイルとシートも残し、出典をたどれるようにする。
                records.append((
                    year,
                    month,
                    prefecture_code,
                    prefecture,
                    int(guest_nights),
                    filename,
                    sheet.title,
                ))

            print(f"読み込み: {year}年{month}月")

    finally:
        wb.close()

# 年・月・都道府県コードの全組み合わせがあるか確認する。
expected_keys = {
    (year, month, f"{code:02d}")
    for year in (2023, 2024, 2025)
    for month in range(1, 13)
    for code in range(1, 48)
}
actual_keys = {(r[0], r[1], r[2]) for r in records}

if len(records) != len(expected_keys) or actual_keys != expected_keys:
    raise ValueError("年・月・都道府県に不足、余分、または重複があります。")

# 同じ都道府県コードに、異なる名前が付いていないか確認する。
for code in range(1, 48):
    names = {r[3] for r in records if r[2] == f"{code:02d}"}
    if len(names) != 1:
        raise ValueError(f"都道府県名が一致しません: {code}: {names}")

con = duckdb.connect()

try:
    # 列の名前と型を指定する。
    # コードは「01」を維持するため、文字列として保存する。
    con.execute("""
        CREATE TABLE accommodation_monthly (
            year INTEGER,
            month INTEGER,
            prefecture_code VARCHAR,
            prefecture VARCHAR,
            guest_nights BIGINT,
            source_file VARCHAR,
            source_sheet VARCHAR
        )
    """)

    # ?に、各行の7つの値が順番に入る。
    # executemanyは、複数行に対して同じSQLを実行する。
    con.executemany(
        "INSERT INTO accommodation_monthly VALUES (?, ?, ?, ?, ?, ?, ?)",
        records,
    )

    print("\n年別の確認（各年564行＝12か月×47都道府県）")
    con.sql("""
        SELECT
            year,
            count(*) AS rows,
            count(DISTINCT month) AS months,
            count(DISTINCT prefecture_code) AS prefectures
        FROM accommodation_monthly
        GROUP BY year
        ORDER BY year
    """).show()

    OUTPUT_FILE.parent.mkdir(parents=True, exist_ok=True)

    # 原本Excelは変更せず、分析用のParquetを保存する。
    # 再実行すると、この出力ファイルを更新する。
    con.sql("""
        SELECT *
        FROM accommodation_monthly
        ORDER BY year, month, prefecture_code
    """).write_parquet(str(OUTPUT_FILE), compression="zstd")

    # 保存したファイルを読み直して、内容の一致を確認する。
    con.read_parquet(str(OUTPUT_FILE)).create_view("saved")
    difference = con.sql("""
        (SELECT * FROM accommodation_monthly EXCEPT ALL SELECT * FROM saved)
        UNION ALL
        (SELECT * FROM saved EXCEPT ALL SELECT * FROM accommodation_monthly)
        LIMIT 1
    """).fetchall()

    if difference:
        raise ValueError("保存前後のデータが一致しません。")

    print(f"保存しました: {OUTPUT_FILE}")
    print(f"{len(records):,}行 / 保存前後の一致を確認済み")

finally:
    con.close()
